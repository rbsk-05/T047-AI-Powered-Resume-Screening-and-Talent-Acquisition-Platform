"""HTTP-level tests for the resume API endpoints.

Uses FastAPI's TestClient with SQLite (configured by conftest.py).
Auth-gated endpoints use dependency overrides to inject a fake user.

Endpoints covered
-----------------
POST   /api/v1/resumes/parse    – no auth, ephemeral
POST   /api/v1/resumes          – candidate auth, persists
GET    /api/v1/resumes          – auth required, role-aware list
GET    /api/v1/resumes/{id}     – auth required, single fetch
DELETE /api/v1/resumes/{id}     – candidate auth, own records only
"""

from __future__ import annotations

import io
import uuid
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.deps import get_current_user, require_candidate
from app.main import app
from app.models.user import User

# ---------------------------------------------------------------------------
# Test fixtures
# ---------------------------------------------------------------------------


def _make_user(role: str = "candidate") -> User:
    u = User()
    u.id = uuid.uuid4()
    u.email = f"{role}@test.com"
    u.role = role
    u.hashed_password = "x"
    return u


FAKE_CANDIDATE = _make_user("candidate")
FAKE_RECRUITER = _make_user("recruiter")


@pytest.fixture()
def client_candidate() -> TestClient:
    """TestClient where all auth resolves to a candidate user."""
    app.dependency_overrides[get_current_user] = lambda: FAKE_CANDIDATE
    app.dependency_overrides[require_candidate] = lambda: FAKE_CANDIDATE
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture()
def client_recruiter() -> TestClient:
    """TestClient where all auth resolves to a recruiter user."""
    app.dependency_overrides[get_current_user] = lambda: FAKE_RECRUITER
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture()
def client_no_auth() -> TestClient:
    """Unauthenticated TestClient (no dependency overrides)."""
    return TestClient(app)


def _docx_bytes(
    name: str = "John Doe",
    email: str = "john@example.com",
    phone: str = "+91 98765 43210",
) -> bytes:
    doc = Document()
    doc.add_paragraph(name)
    doc.add_paragraph(f"{email} | {phone}")
    doc.add_paragraph("linkedin.com/in/johndoe")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python, FastAPI, AWS, Docker, PostgreSQL")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Software Engineer at TechCorp | Jan 2021 - Dec 2023")
    doc.add_paragraph("Education")
    doc.add_paragraph("B.Tech Computer Science")
    doc.add_paragraph("Projects")
    doc.add_paragraph("Cloud E-commerce Platform")
    doc.add_paragraph("Certifications")
    doc.add_paragraph("AWS Cloud Practitioner")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# POST /api/v1/resumes/parse  — ephemeral, no auth
# ---------------------------------------------------------------------------


class TestParseEndpoint:
    def test_valid_docx_returns_profile(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.post(
            "/api/v1/resumes/parse",
            files={"file": ("resume.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["email"] == "john@example.com"
        assert isinstance(data["skills"], list)
        assert isinstance(data["experience_entries"], list)
        # candidate_id should be null for ephemeral parse
        assert data["candidate_id"] is None

    def test_missing_file_returns_422(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.post("/api/v1/resumes/parse")
        assert response.status_code == 422

    def test_unsupported_extension_returns_400(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.post(
            "/api/v1/resumes/parse",
            files={"file": ("resume.txt", b"some text content", "text/plain")},
        )
        assert response.status_code == 400
        assert "Unsupported file type" in response.json()["detail"]

    def test_corrupted_pdf_returns_400(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.post(
            "/api/v1/resumes/parse",
            files={"file": ("resume.pdf", b"not a real pdf", "application/pdf")},
        )
        assert response.status_code == 400

    def test_oversized_file_returns_413(self, client_no_auth: TestClient) -> None:
        big_content = b"x" * (11 * 1024 * 1024)  # 11 MB
        response = client_no_auth.post(
            "/api/v1/resumes/parse",
            files={"file": ("resume.docx", big_content, "application/octet-stream")},
        )
        assert response.status_code == 413

    def test_non_resume_document_returns_400(self, client_no_auth: TestClient) -> None:
        doc = Document()
        doc.add_paragraph("Dear Sir, please find attached my application.")
        doc.add_paragraph("Yours sincerely, Someone")
        buf = io.BytesIO()
        doc.save(buf)
        response = client_no_auth.post(
            "/api/v1/resumes/parse",
            files={"file": ("letter.docx", buf.getvalue(), "application/octet-stream")},
        )
        assert response.status_code == 400


# ---------------------------------------------------------------------------
# POST /api/v1/resumes  — persist (candidate auth)
# ---------------------------------------------------------------------------


class TestCreateCandidateEndpoint:
    def test_valid_upload_returns_201_with_id(
        self, client_candidate: TestClient
    ) -> None:
        response = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 201
        data = response.json()
        assert "id" in data
        assert data["profile"]["email"] == "john@example.com"
        # candidate_id should be populated after persist
        assert data["profile"]["candidate_id"] is not None
        assert data["profile"]["candidate_id"] == data["id"]

    def test_upload_requires_auth(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 401

    def test_invalid_file_type_returns_400(
        self, client_candidate: TestClient
    ) -> None:
        response = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.txt", b"not a resume", "text/plain")},
        )
        assert response.status_code == 400


# ---------------------------------------------------------------------------
# GET /api/v1/resumes  — list
# ---------------------------------------------------------------------------


class TestListCandidatesEndpoint:
    def test_list_requires_auth(self, client_no_auth: TestClient) -> None:
        response = client_no_auth.get("/api/v1/resumes")
        assert response.status_code == 401

    def test_candidate_can_list(self, client_candidate: TestClient) -> None:
        response = client_candidate.get("/api/v1/resumes")
        assert response.status_code == 200
        assert isinstance(response.json(), list)

    def test_recruiter_can_list(self, client_recruiter: TestClient) -> None:
        response = client_recruiter.get("/api/v1/resumes")
        assert response.status_code == 200
        assert isinstance(response.json(), list)


# ---------------------------------------------------------------------------
# GET /api/v1/resumes/{id}  — single fetch
# ---------------------------------------------------------------------------


class TestGetCandidateEndpoint:
    def test_nonexistent_id_returns_404(self, client_recruiter: TestClient) -> None:
        fake_id = uuid.uuid4()
        response = client_recruiter.get(f"/api/v1/resumes/{fake_id}")
        assert response.status_code == 404

    def test_requires_auth(self, client_no_auth: TestClient) -> None:
        fake_id = uuid.uuid4()
        response = client_no_auth.get(f"/api/v1/resumes/{fake_id}")
        assert response.status_code == 401

    def test_candidate_can_fetch_own_record(
        self, client_candidate: TestClient
    ) -> None:
        # Upload first
        upload = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert upload.status_code == 201
        candidate_id = upload.json()["id"]

        # Fetch it back
        response = client_candidate.get(f"/api/v1/resumes/{candidate_id}")
        assert response.status_code == 200
        assert response.json()["id"] == candidate_id

    def test_recruiter_can_fetch_any_record(
        self, client_candidate: TestClient, client_recruiter: TestClient
    ) -> None:
        upload = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert upload.status_code == 201
        candidate_id = upload.json()["id"]

        response = client_recruiter.get(f"/api/v1/resumes/{candidate_id}")
        assert response.status_code == 200

    def test_candidate_cannot_fetch_others_record(
        self, client_candidate: TestClient
    ) -> None:
        """A candidate accessing a record owned by a different user gets 403."""
        from app.api.deps import require_candidate

        other_user = _make_user("candidate")

        # Upload as FAKE_CANDIDATE (the default override)
        upload = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert upload.status_code == 201
        candidate_id = upload.json()["id"]

        # Now switch to a different candidate user
        app.dependency_overrides[get_current_user] = lambda: other_user
        app.dependency_overrides[require_candidate] = lambda: other_user
        other_client = TestClient(app)
        response = other_client.get(f"/api/v1/resumes/{candidate_id}")
        # Reset overrides
        app.dependency_overrides[get_current_user] = lambda: FAKE_CANDIDATE
        app.dependency_overrides[require_candidate] = lambda: FAKE_CANDIDATE
        assert response.status_code == 403


# ---------------------------------------------------------------------------
# DELETE /api/v1/resumes/{id}
# ---------------------------------------------------------------------------


class TestDeleteCandidateEndpoint:
    def test_candidate_can_delete_own_record(
        self, client_candidate: TestClient
    ) -> None:
        upload = client_candidate.post(
            "/api/v1/resumes",
            files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")},
        )
        assert upload.status_code == 201
        candidate_id = upload.json()["id"]

        delete_resp = client_candidate.delete(f"/api/v1/resumes/{candidate_id}")
        assert delete_resp.status_code == 204

        # Verify it's gone
        fetch_resp = client_candidate.get(f"/api/v1/resumes/{candidate_id}")
        assert fetch_resp.status_code == 404

    def test_nonexistent_id_returns_404(self, client_candidate: TestClient) -> None:
        response = client_candidate.delete(f"/api/v1/resumes/{uuid.uuid4()}")
        assert response.status_code == 404

    def test_recruiter_cannot_delete(self, client_recruiter: TestClient) -> None:
        """DELETE is restricted to candidate role."""
        response = client_recruiter.delete(f"/api/v1/resumes/{uuid.uuid4()}")
        assert response.status_code == 403
